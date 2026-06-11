import os
import io
import json
import time
import pdfplumber
import streamlit as st
from dotenv import load_dotenv
from openai import OpenAI

st.set_page_config(
    page_title="Extractor Fiscal",
    page_icon=":money_with_wings:",
    layout="wide",
)

try:
    api_key = st.secrets["OPENAI_API_KEY"]
except (KeyError, FileNotFoundError):
    load_dotenv()
    api_key = os.getenv("OPENAI_API_KEY")

if not api_key:
    st.error("La llave API no está configurada.")
    st.stop()

assert isinstance(api_key, str)
client = OpenAI(api_key=api_key)

MAX_ARCHIVOS = 5
MAX_SIZE_MB = 2

def _leer_txt(data: bytes) -> str:
    for enc in ("utf-8", "latin-1", "cp1252"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _leer_pdf(data: bytes) -> str:
  
    paginas = []
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        for page in pdf.pages:
            texto = page.extract_text()
            if texto:
                paginas.append(texto)
    if not paginas:
        raise ValueError("El PDF no contiene texto extraíble (puede ser una imagen escaneada).")
    return "\n".join(paginas)


def _leer_docx(data: bytes) -> str:
    from docx import Document
    doc = Document(io.BytesIO(data))
    lineas = [p.text for p in doc.paragraphs if p.text.strip()]
    if not lineas:
        raise ValueError("El documento Word no contiene párrafos con texto.")
    return "\n".join(lineas)


def _leer_xlsx(data: bytes) -> str:
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    lineas = []
    for sheet in wb.worksheets:
        for row in sheet.iter_rows(values_only=True):
            celda = "\t".join("" if c is None else str(c) for c in row)
            if celda.strip():
                lineas.append(celda)
    if not lineas:
        raise ValueError("La hoja de cálculo no contiene datos.")
    return "\n".join(lineas)


_LECTORES = {
    ".txt": _leer_txt,
    ".pdf": _leer_pdf,
    ".docx": _leer_docx,
    ".xlsx": _leer_xlsx,
}


def leer_archivo(archivo) -> tuple[str | None, str | None]:
    """Retorna (texto, mensaje_error). Uno de los dos siempre es None."""
    ext = os.path.splitext(archivo.name)[1].lower()
    lector = _LECTORES.get(ext)
    if lector is None:
        return None, f"Formato no soportado: '{ext}'"
    try:
        return lector(archivo.getvalue()), None
    except Exception as e:
        return None, f"{type(e).__name__}: {e}"



_SYSTEM_PROMPT = (
    "You are a fiscal data extractor. "
    "You respond ONLY with valid JSON, no additional text. "
    "You handle documents in any language (Spanish, English, or mixed)."
)

_USER_PROMPT = """\
Extract the following fields from the text below and respond ONLY with the JSON.
The document may be in Spanish, English, or a mix of both — detect the language automatically.

REQUIRED SCHEMA:
{{
  "nombre_cliente": "full client name (or 'Name', 'Client', 'Customer' in English docs)",
  "monto": "digits only, no $ or commas (e.g. 47850)",
  "fecha": "YYYY-MM-DD format",
  "tipo_solicitud": "VENTA/SALE | QUEJA/COMPLAINT | FACTURA/INVOICE"
}}

Use "NO_ENCONTRADO" for any field not present in the text.
Normalize tipo_solicitud to one of: VENTA, QUEJA, FACTURA — regardless of the source language.

TEXT TO ANALYZE:
{texto}
"""


def extraer_datos_con_ia(texto: str) -> str:
    response = client.chat.completions.create(
        model="gpt-4o-mini",
        messages=[
            {"role": "system", "content": _SYSTEM_PROMPT},
            {"role": "user", "content": _USER_PROMPT.format(texto=texto)},
        ],
        response_format={"type": "json_object"},
        temperature=0,
    )
    return response.choices[0].message.content or ""


def parsear_json(respuesta: str) -> dict | None:
    limpia = (
        respuesta.strip()
        .removeprefix("```json")
        .removeprefix("```")
        .removesuffix("```")
        .strip()
    )
    try:
        return json.loads(limpia)
    except json.JSONDecodeError:
        return None


def tiene_datos_utiles(datos: dict) -> bool:
    return any(v and v != "NO_ENCONTRADO" for v in datos.values())



def procesar_archivo(archivo) -> dict:
    texto, error_lectura = leer_archivo(archivo)

    if error_lectura:
        return {"datos": None, "estado": "error_lectura", "mensaje": error_lectura}

    if not texto or not texto.strip():
        return {"datos": None, "estado": "sin_datos", "mensaje": "Archivo vacío o sin texto legible."}

    placeholder = st.empty()

    for intento in range(1, 4):
        placeholder.info(f"Consultando IA — intento {intento}/3…")
        try:
            respuesta = extraer_datos_con_ia(texto)
            datos = parsear_json(respuesta)
        except Exception as e:
            if intento == 3:
                placeholder.empty()
                return {"datos": None, "estado": "error_ia", "mensaje": str(e)}
            time.sleep(1)
            continue

        if datos:
            if tiene_datos_utiles(datos):
                placeholder.success(f"Datos extraídos en intento {intento}.")
                return {"datos": datos, "estado": "exitoso", "mensaje": f"Intento {intento}"}
            else:
                placeholder.empty()
                return {"datos": None, "estado": "sin_datos", "mensaje": "Sin información relevante."}

        if intento < 3:
            time.sleep(1)

    placeholder.empty()
    return {
        "datos": None,
        "estado": "error_procesamiento",
        "mensaje": "No se pudo obtener JSON válido tras 3 intentos.",
    }



st.title("Extractor Fiscal — Text-to-Structured-Data")
st.markdown("---")

st.markdown(f"""
### Objetivo
Convertir textos desordenados en datos estructurados (JSON) para sistemas administrativos o fiscales.

### Formatos soportados
- `.txt` — texto plano
- `.pdf` — documentos PDF (con texto seleccionable)
- `.docx` — documentos Word
- `.xlsx` — hojas de cálculo Excel

> **Demo:** máx {MAX_ARCHIVOS} archivos · {MAX_SIZE_MB} MB por archivo
""")

st.markdown("---")
st.subheader("Subir archivos")

archivos_subidos = st.file_uploader(
    "Arrastra o selecciona uno o más archivos",
    type=["txt", "pdf", "docx", "xlsx"],
    accept_multiple_files=True,
)

st.markdown("---")

if archivos_subidos:
    if len(archivos_subidos) > MAX_ARCHIVOS:
        st.error(f"Máximo {MAX_ARCHIVOS} archivos por sesión.")
        st.stop()

    for archivo in archivos_subidos:
        size_mb = archivo.size / (1024 * 1024)
        if size_mb > MAX_SIZE_MB:
            st.error(f"`{archivo.name}` pesa {size_mb:.1f} MB (máx {MAX_SIZE_MB} MB).")
            st.stop()

    st.success(f"{len(archivos_subidos)} archivo(s) cargado(s).")

    if st.button("Procesar archivos", type="primary"):
        resultados = []
        sin_datos = 0

        barra = st.progress(0)
        estado_texto = st.empty()

        for i, archivo in enumerate(archivos_subidos):
            barra.progress((i + 1) / len(archivos_subidos))
            estado_texto.text(f"Procesando {i + 1}/{len(archivos_subidos)}: {archivo.name}")

            resultado = procesar_archivo(archivo)

            if resultado["estado"] == "sin_datos":
                sin_datos += 1
            else:
                resultados.append({"archivo": archivo.name, **resultado})

        barra.empty()
        estado_texto.empty()

        if sin_datos:
            st.info(f"{sin_datos} archivo(s) sin información relevante.")

        if resultados:
            st.markdown("---")
            st.subheader("Resultados")

            for r in resultados:
                icono =  "" if r["estado"] == "exitoso" else ""
                with st.expander(f"{icono} {r['archivo']}", expanded=(r["estado"] == "exitoso")):
                    if r["estado"] == "exitoso":
                        st.json(r["datos"])
                    else:
                        st.error(r["mensaje"])

            exitosos = [r for r in resultados if r["estado"] == "exitoso"]
            if exitosos:
                resumen = {
                    "total_archivos_procesados": len(archivos_subidos),
                    "archivos_exitosos": len(exitosos),
                    "archivos_sin_datos": sin_datos,
                    "archivos_con_errores": len(resultados) - len(exitosos),
                    "resultados": exitosos,
                }
                st.markdown("---")
                st.download_button(
                    label="Descargar resultados (JSON)",
                    data=json.dumps(resumen, indent=2, ensure_ascii=False),
                    file_name="resultados_extraccion.json",
                    mime="application/json",
                )
            else:
                st.warning("No se pudieron extraer datos válidos.")
        else:
            st.info("No se encontraron datos relevantes en los archivos.")

else:
    st.info("Sube uno o más archivos para comenzar.")

st.markdown("---")
st.markdown(
    "<div style='text-align:center'><p style='color:gray;'>Brandon Enrique Eroza Torres</p></div>",
    unsafe_allow_html=True,
)
